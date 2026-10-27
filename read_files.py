import sys
print("Starting file extraction...", flush=True)

try:
    from docx import Document
    print("docx imported", flush=True)

    doc = Document(r'e:\CaliboAiAcad\MCU2\CALIBO_USECASE-2\KLU USECASE-2\MUC02_KLU_Credit_PreScreening.docx')
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    with open(r'e:\CaliboAiAcad\MCU2\doc_out.txt', 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print("doc_out.txt written", flush=True)

    doc2 = Document(r'e:\CaliboAiAcad\MCU2\CALIBO_USECASE-2\KLU USECASE-2\MUC02_KLU_Trainer_Notes.docx')
    lines2 = [p.text for p in doc2.paragraphs if p.text.strip()]
    with open(r'e:\CaliboAiAcad\MCU2\trainer_notes.txt', 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines2))
    print("trainer_notes.txt written", flush=True)

except Exception as e:
    print(f"DOCX error: {e}", flush=True)

try:
    from pptx import Presentation
    import os
    print("pptx imported", flush=True)
    base = r'e:\CaliboAiAcad\MCU2\CALIBO_USECASE-2\KLU USECASE-2'
    for fname, outname in [('MUC02_KLU_Expected_Outcomes.pptx', 'outcomes.txt'),
                           ('MUC02_KLU_Trainer_Briefing.pptx', 'briefing.txt')]:
        prs = Presentation(os.path.join(base, fname))
        lines = []
        for i, slide in enumerate(prs.slides):
            lines.append(f'=== SLIDE {i+1} ===')
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text.strip():
                    lines.append(shape.text.strip())
        with open(r'e:\CaliboAiAcad\MCU2\\' + outname, 'w', encoding='utf-8') as f:
            f.write('\n'.join(lines))
        print(f"{outname} written", flush=True)
except Exception as e:
    print(f"PPTX error: {e}", flush=True)

try:
    import pandas as pd
    print("pandas imported", flush=True)
    df = pd.read_csv(r'e:\CaliboAiAcad\MCU2\CALIBO_USECASE-2\KLU USECASE-2\MUC02_Credit_PreScreening_Dataset 1.csv')
    lines = []
    lines.append('SHAPE: ' + str(df.shape))
    lines.append('COLUMNS:\n' + '\n'.join(df.columns.tolist()))
    lines.append('DTYPES:\n' + str(df.dtypes))
    lines.append('MISSING VALUES:\n' + str(df.isnull().sum()))
    lines.append('HEAD (3 rows):\n' + df.head(3).to_string())
    lines.append('TARGET (last col) value counts:\n' + str(df.iloc[:, -1].value_counts()))
    lines.append('DESCRIBE:\n' + str(df.describe().to_string()))
    with open(r'e:\CaliboAiAcad\MCU2\csv_info.txt', 'w', encoding='utf-8') as f:
        f.write('\n\n'.join(lines))
    print("csv_info.txt written", flush=True)
except Exception as e:
    print(f"CSV error: {e}", flush=True)

print("ALL DONE", flush=True)
